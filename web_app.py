import shutil
import tempfile
import uuid
from pathlib import Path

import openpyxl
from flask import Flask, jsonify, render_template, request, send_file
from werkzeug.utils import secure_filename

import aggregate
from app import parse_phieu, write_to_tong

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "web_uploads"
OUTPUT_DIR = BASE_DIR / "web_outputs"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024


class SilentMessageBox:
    @staticmethod
    def showinfo(*args, **kwargs):
        return None

    showwarning = showinfo
    showerror = showinfo


def make_job_dir():
    return Path(tempfile.mkdtemp(prefix="pcgd_", dir=UPLOAD_DIR))


def save_upload(file_storage, folder):
    filename = secure_filename(file_storage.filename or "file.xlsx")
    if not filename.lower().endswith(".xlsx") or filename.startswith("~$"):
        raise ValueError("Chỉ hỗ trợ file Excel .xlsx")
    path = folder / f"{uuid.uuid4().hex}_{filename}"
    file_storage.save(path)
    return path


def collect_xlsx_files(file_storages):
    files = []
    seen = set()
    for file_storage in file_storages or []:
        if not file_storage or not getattr(file_storage, "filename", None):
            continue
        filename = str(file_storage.filename).replace('\\', '/')
        safe_name = secure_filename(filename)
        if not safe_name.lower().endswith(".xlsx") or safe_name.startswith("~$"):
            continue
        key = filename
        if key in seen:
            continue
        seen.add(key)
        files.append(file_storage)
    return files


def copy_result(source_path, download_name):
    output_path = OUTPUT_DIR / f"{uuid.uuid4().hex[:8]}_{download_name}"
    shutil.copy2(source_path, output_path)
    return f"/download/{output_path.name}"


def resolve_template_path(template_name):
    candidates = [
        BASE_DIR / template_name,
        BASE_DIR / "templates" / template_name,
    ]
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def get_template_path(uploaded_file, fallback_name):
    if uploaded_file and getattr(uploaded_file, "filename", None):
        return uploaded_file
    default_path = resolve_template_path(fallback_name)
    if default_path is None:
        raise ValueError(f"Không tìm thấy file mẫu '{fallback_name}' trong thư mục dự án.")
    return default_path


def build_tong(phieu_files, tong_template):
    output_path = OUTPUT_DIR / f"KetQua_Tong_{uuid.uuid4().hex[:8]}.xlsx"
    shutil.copy2(tong_template, output_path)
    workbook = openpyxl.load_workbook(output_path)
    if workbook.active.max_row > 4:
        workbook.active.delete_rows(5, workbook.active.max_row - 4)

    success = 0
    errors = []
    for source in phieu_files:
        general_info, people, parse_error = parse_phieu(str(source))
        if not general_info or not people:
            errors.append({"file": source.name, "reason": parse_error or "Không tìm thấy dữ liệu người"})
            continue
        try:
            write_to_tong(general_info, people, str(output_path), workbook=workbook)
            success += 1
        except Exception as error:
            errors.append({"file": source.name, "reason": str(error)})
    workbook.save(output_path)
    workbook.close()
    return output_path, success, errors


def required_upload(name):
    file_storage = request.files.get(name)
    if not file_storage or not file_storage.filename:
        raise ValueError(f"Vui lòng chọn file {name}.")
    return file_storage


def run_report(operation, source_name, template_names):
    job_dir = make_job_dir()
    source = save_upload(required_upload(source_name), job_dir)
    default_template_map = {
        "m1_template": "M1.xlsx",
        "m2_template": "M2.xlsx",
        "spc1_template": "SPC1.xlsx",
        "spcth_template": "SPCTH.xlsx",
        "m1th_template": "M1TH.xlsx",
        "xmth_template": "XMTH.xlsx",
        "pl_template": "PL.xlsx",
    }
    templates = []
    for template_field in template_names:
        uploaded = request.files.get(template_field)
        fallback_name = default_template_map.get(template_field)
        resolved = get_template_path(uploaded, fallback_name)
        if hasattr(resolved, "filename") and getattr(resolved, "filename", None):
            templates.append(save_upload(resolved, job_dir))
        else:
            templates.append(resolved)
    aggregate.messagebox = SilentMessageBox

    if operation == "m1-m2":
        aggregate.aggregate_m1_m2(str(source), str(templates[0]), str(templates[1]), "")
        outputs = [(job_dir / "KetQua_M1.xlsx", "KetQua_M1.xlsx"), (job_dir / "KetQua_M2.xlsx", "KetQua_M2.xlsx")]
    elif operation == "spc1":
        aggregate.aggregate_spc1(str(source), str(templates[0]))
        outputs = [(job_dir / "KetQua_SPC1.xlsx", "KetQua_SPC1.xlsx")]
    elif operation == "primary":
        aggregate.aggregate_primary(str(source), *(str(template) for template in templates))
        outputs = [(job_dir / name, name) for name in ("KetQua_SPCTH.xlsx", "KetQua_M1TH.xlsx", "KetQua_XMTH.xlsx")]
    elif operation == "appendix":
        aggregate.export_appendix(str(source), str(templates[0]))
        outputs = [(job_dir / "KetQua_PhụLục.xlsx", "KetQua_PhụLục.xlsx")]
    else:
        raise ValueError("Thao tác không được hỗ trợ.")

    downloads = []
    for output, name in outputs:
        if not output.exists():
            raise RuntimeError(f"Không tạo được {name}.")
        downloads.append({"name": name, "url": copy_result(output, name)})
    return downloads


def review_tong(source_file):
    from app import App

    workbook = openpyxl.load_workbook(source_file, data_only=True)
    sheet = workbook.active
    records = []
    for row_number in range(6, sheet.max_row + 1):
        values = [sheet.cell(row=row_number, column=column).value for column in range(1, 48)]
        if all(App._is_missing(value) for value in values):
            continue
        try:
            birth_year = int(float(values[5]))
        except (TypeError, ValueError):
            birth_year = None
        person = " ".join(str(values[index]).strip() for index in (1, 2) if not App._is_missing(values[index]))
        village, so_phieu = values[12], values[13]
        missing = []
        if birth_year is None:
            missing.append(("M1, M2", "Năm sinh"))
        if App._is_missing(village):
            missing.append(("M2", "Địa chỉ/thôn bản"))
        if App._is_missing(values[7]):
            missing.append(("M1", "Dân tộc"))
        studying = not App._is_missing(values[17]) and App._is_missing(values[24]) and App._is_missing(values[26])
        age_group = bool(birth_year and 2008 <= birth_year <= 2014)
        if studying and App._is_missing(values[16]):
            missing.append(("M1", "Khối học"))
        if studying and age_group and App._is_missing(values[16]):
            missing.append(("M2", "Khối học"))
        if age_group:
            if App._is_missing(values[19]):
                missing.append(("M2", "Bậc tốt nghiệp"))
            if App._is_missing(values[21]):
                missing.append(("M2", "Năm tốt nghiệp"))
        if not App._is_missing(values[21]) and not App._is_academic_year(values[21]):
            missing.append(("M2", "Năm tốt nghiệp không đúng dạng YYYY-YYYY"))
        mismatch = App._graduation_class_mismatch(values[17], values[16], values[19])
        if mismatch:
            missing.append(("M1, M2", f"Cấp/lớp không phù hợp: lớp {mismatch[1]} nhưng bậc tốt nghiệp {mismatch[0]}"))
        for class_index, year_index, label in ((24, 25, "học xong"), (26, 27, "bỏ học")):
            class_value, year_value = values[class_index], values[year_index]
            if not App._is_missing(class_value) and App._is_missing(year_value):
                missing.append(("M1, M2", f"Năm {label}"))
            if not App._is_missing(year_value) and not App._is_academic_year(year_value):
                missing.append(("M1, M2", f"Năm {label} không đúng dạng YYYY-YYYY"))
            if not App._is_missing(year_value) and App._is_missing(class_value):
                missing.append(("M1, M2", f"Lớp {label}"))
        for report, field in missing:
            records.append({"row": row_number, "so_phieu": "" if App._is_missing(so_phieu) else str(so_phieu), "person": person, "birth_year": birth_year or "", "village": "" if App._is_missing(village) else str(village), "report": report, "field": field})
    workbook.close()
    return records


def export_review_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thông tin cần bổ sung"
    ws.append([
        "STT", "Dòng trong file Tổng", "Số phiếu", "Họ tên", "Năm sinh",
        "Thôn/bản", "Biểu", "Trường còn thiếu", "Nguyên nhân"
    ])
    for index, record in enumerate(records, start=1):
        ws.append([
            index,
            record.get("row", ""),
            record.get("so_phieu", ""),
            record.get("person", ""),
            record.get("birth_year", ""),
            record.get("village", ""),
            record.get("report", ""),
            "",
            record.get("field", "")
        ])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 25
    ws.column_dimensions["I"].width = 60
    output_path = OUTPUT_DIR / f"DanhSachCanBoSung_{uuid.uuid4().hex[:8]}.xlsx"
    wb.save(output_path)
    wb.close()
    return output_path


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/api/tong-hop")
def aggregate_forms():
    phieu_files = collect_xlsx_files(request.files.getlist("phieu_files"))
    if not phieu_files:
        return jsonify({"error": "Vui lòng chọn ít nhất một file phiếu .xlsx."}), 400
    job_dir = make_job_dir()
    try:
        forms = [save_upload(file, job_dir) for file in phieu_files if file.filename]
        template_upload = request.files.get("tong_template")
        template_path = get_template_path(template_upload, "Tong.xlsx")
        if hasattr(template_path, "filename") and getattr(template_path, "filename", None):
            template = save_upload(template_path, job_dir)
        else:
            template = str(template_path)
        output, success, errors = build_tong(forms, template)
        return jsonify({"message": f"Đã tổng hợp {success} phiếu.", "success": success, "errors": errors, "download_url": f"/download/{output.name}"})
    except Exception as error:
        return jsonify({"error": f"Không thể tổng hợp: {error}"}), 500


@app.post("/api/bao-cao/<operation>")
def report_operation(operation):
    configs = {
        "m1-m2": ("tong_file", ["m1_template", "m2_template"]),
        "spc1": ("tong_file", ["spc1_template"]),
        "primary": ("tong_file", ["spcth_template", "m1th_template", "xmth_template"]),
        "appendix": ("tong_file", ["pl_template"]),
    }
    if operation not in configs:
        return jsonify({"error": "Thao tác không được hỗ trợ."}), 404
    try:
        source_name, template_names = configs[operation]
        downloads = run_report(operation, source_name, template_names)
        return jsonify({"message": "Đã tạo báo cáo thành công.", "downloads": downloads})
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except Exception as error:
        return jsonify({"error": f"Không thể tạo báo cáo: {error}"}), 500


@app.post("/api/ra-soat")
def review_operation():
    job_dir = make_job_dir()
    try:
        source = save_upload(required_upload("tong_file"), job_dir)
        records = review_tong(source)
        download_url = None
        if records:
            output_path = export_review_excel(records)
            download_url = f"/download/{output_path.name}"
        return jsonify({
            "count": len(records),
            "records": records,
            "message": f"Đã rà soát, phát hiện {len(records)} mục cần kiểm tra.",
            "download_url": download_url,
        })
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except Exception as error:
        return jsonify({"error": f"Không thể rà soát: {error}"}), 500


@app.post("/api/ra-soat/export")
def review_export():
    job_dir = make_job_dir()
    try:
        source = save_upload(required_upload("tong_file"), job_dir)
        records = review_tong(source)
        if not records:
            return jsonify({"message": "Không có dữ liệu cần bổ sung để xuất.", "download_url": None}), 200
        output_path = export_review_excel(records)
        return jsonify({"message": "Đã xuất file rà soát thành công.", "download_url": f"/download/{output_path.name}"})
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except Exception as error:
        return jsonify({"error": f"Không thể xuất file rà soát: {error}"}), 500


@app.get("/download/<filename>")
def download(filename):
    requested_name = Path(filename).name
    path = OUTPUT_DIR / requested_name
    if requested_name != filename or not path.exists() or path.parent != OUTPUT_DIR:
        return jsonify({"error": "Không tìm thấy file kết quả."}), 404
    return send_file(path, as_attachment=True, download_name=requested_name.split("_", 1)[-1])


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
